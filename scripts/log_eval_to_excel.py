#!/usr/bin/env python
"""Log one eval run into the experiment Excel workbook (default docs/il_ablation_future_feat.xlsx).

Usage:
    python scripts/log_eval_to_excel.py \
        --eval-log log/eval/eval_all_20260925_111916.log \
        --name "B_baseline" \
        --checkpoint checkpoints_vla/xxx/checkpoint_epoch_40.pt \
        --note "future_feat=off, next_chunk=off"

Excel structure:
    Experiments        one row per experiment (upsert by experiment name)
    Data               long table: experiment x task -> success rate (source of truth, upsert)
    TaskSuccessMatrix  rebuilt from the Data sheet: tasks as rows, experiments as columns, last row = mean
Eval log format: task,ok,total,rate lines + a MEAN_SUCCESS,<rate> line
(output of scripts/eval_metaworld_parallel.py).
"""
import argparse
import os
import re
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="4472C4")
HEADER_FONT = Font(bold=True, color="FFFFFF")
MEAN_FILL = PatternFill("solid", fgColor="D9E1F2")


def parse_eval_log(path):
    rates = {}
    mean = None
    for line in open(path):
        m = re.match(r"(\S+-v3),(\d+),(\d+),([\d.]+)", line)
        if m:
            rates[m.group(1)] = (int(m.group(2)), int(m.group(3)), 100.0 * float(m.group(4)))
        m = re.search(r"MEAN_SUCCESS,([\d.]+)", line)
        if m:
            mean = 100.0 * float(m.group(1))
    if not rates:
        raise ValueError(f"no task,ok,total,rate lines parsed from {path}")
    if mean is None:
        mean = sum(v[2] for v in rates.values()) / len(rates)
    return rates, mean


def style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(1, c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")


def upsert_row(ws, key_col, key, values):
    """Find key in column key_col; overwrite that row if found, else append. values = full row."""
    for row in range(2, ws.max_row + 1):
        if ws.cell(row, key_col).value == key:
            for c, v in enumerate(values, 1):
                ws.cell(row, c, v)
            return
    ws.append(values)


def main():
    ap = argparse.ArgumentParser(description="Log eval results into the experiment Excel workbook")
    ap.add_argument("--eval-log", required=True)
    ap.add_argument("--name", required=True, help="experiment name (upsert key)")
    ap.add_argument("--checkpoint", default="")
    ap.add_argument("--note", default="", help="notes such as config differences")
    ap.add_argument("--excel", default="docs/il_ablation_future_feat.xlsx")
    args = ap.parse_args()

    rates, mean = parse_eval_log(args.eval_log)

    if os.path.exists(args.excel):
        wb = load_workbook(args.excel)
    else:
        wb = Workbook()
        wb.remove(wb.active)

    # ---------- experiment log sheet ----------
    if "Experiments" in wb.sheetnames:
        ws_exp = wb["Experiments"]
    else:
        ws_exp = wb.create_sheet("Experiments", 0)
        ws_exp.append(["Experiment", "Date", "Checkpoint", "Note", "Mean Success(%)", "Eval Log"])
        style_header(ws_exp, 6)
    upsert_row(ws_exp, 1, args.name,
               [args.name, datetime.now().strftime("%Y-%m-%d %H:%M"),
                args.checkpoint, args.note, round(mean, 2), args.eval_log])

    # ---------- data sheet (long table) ----------
    if "Data" in wb.sheetnames:
        ws_data = wb["Data"]
    else:
        ws_data = wb.create_sheet("Data")
        ws_data.append(["Experiment", "Task", "Success", "Total", "Rate(%)"])
        style_header(ws_data, 5)
    for task in sorted(rates):
        ok, total, rate = rates[task]
        # try updating an existing (experiment, task) row first
        updated = False
        for row in range(2, ws_data.max_row + 1):
            if ws_data.cell(row, 1).value == args.name and ws_data.cell(row, 2).value == task:
                ws_data.cell(row, 3, ok)
                ws_data.cell(row, 4, total)
                ws_data.cell(row, 5, round(rate, 1))
                updated = True
                break
        if not updated:
            ws_data.append([args.name, task, ok, total, round(rate, 1)])

    # ---------- task success-rate matrix (rebuilt) ----------
    # experiment order follows the row order in the Experiments sheet
    experiments = []
    for row in range(2, ws_exp.max_row + 1):
        v = ws_exp.cell(row, 1).value
        if v:
            experiments.append(v)
    data = {}  # (exp, task) -> rate
    tasks = set()
    for row in range(2, ws_data.max_row + 1):
        exp, task = ws_data.cell(row, 1).value, ws_data.cell(row, 2).value
        if exp and task:
            data[(exp, task)] = ws_data.cell(row, 5).value
            tasks.add(task)

    if "TaskSuccessMatrix" in wb.sheetnames:
        del wb["TaskSuccessMatrix"]
    ws_m = wb.create_sheet("TaskSuccessMatrix", 1)
    ws_m.append(["Task"] + experiments)
    style_header(ws_m, 1 + len(experiments))
    for task in sorted(tasks):
        ws_m.append([task] + [data.get((e, task)) for e in experiments])
    mean_row = ws_m.max_row + 1
    ws_m.cell(mean_row, 1, "Mean").font = Font(bold=True)
    for c in range(2, 2 + len(experiments)):
        vals = [ws_m.cell(r, c).value for r in range(2, mean_row)]
        vals = [v for v in vals if isinstance(v, (int, float))]
        cell = ws_m.cell(mean_row, c, round(sum(vals) / len(vals), 2) if vals else None)
        cell.font = Font(bold=True)
        cell.fill = MEAN_FILL
    ws_m.cell(mean_row, 1).fill = MEAN_FILL
    ws_m.column_dimensions["A"].width = 28
    for c in range(2, 2 + len(experiments)):
        ws_m.column_dimensions[get_column_letter(c)].width = max(16, len(experiments[c - 2]) + 4)
    ws_m.freeze_panes = "B2"

    # column widths: experiment log sheet
    for col, w in zip("ABCDEF", [24, 16, 70, 50, 14, 50]):
        ws_exp.column_dimensions[col].width = w
    for col, w in zip("ABCDE", [24, 28, 8, 8, 10]):
        ws_data.column_dimensions[col].width = w

    wb.save(args.excel)
    print(f"[log_eval_to_excel] {args.name}: mean={mean:.2f}% over {len(rates)} tasks -> {args.excel}")


if __name__ == "__main__":
    main()
